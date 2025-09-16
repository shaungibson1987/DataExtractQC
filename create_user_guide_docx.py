from docx import Document
from docx.shared import Pt

def create_user_guide_docx(path):
    doc = Document()
    doc.add_heading('DataExtractQC User Guide', 0)

    doc.add_paragraph(
        "Welcome to DataExtractQC! This guide will help you use the program to extract and organize survey data from Excel files, even if you have no technical background."
    )

    doc.add_heading('What is DataExtractQC?', level=1)
    doc.add_paragraph(
        "DataExtractQC is a tool for extracting, cleaning, and organizing survey data from Excel files. It is designed for people who work with survey results and need to quickly create clean, organized files for analysis or reporting."
    )

    doc.add_heading('Key Features', level=1)
    features = [
        "Simple graphical interface (no coding required)",
        "Choose which columns to keep using an 'include.txt' file",
        "Optional: Automatically find and add open-ended (text) columns",
        "Optional: Search for specific words and highlight matches",
        "Optional: Check for duplicate postcode/YOB pairs",
        "Optional: Length check to flag unusually long cells (multiplier adjustable)",
        "All advanced checks are off by default—enable only what you need",
        "Output split by language and overall",
        "Detailed logging and error reporting"
    ]
    for feat in features:
        doc.add_paragraph(feat, style='List Bullet')

    doc.add_heading('How to Use DataExtractQC', level=1)
    steps = [
        ("Open the program", [
            "Double-click the data_extract_gui.exe file (or run python data_extract_gui.py if you use Python)."
        ]),
        ("Select your files and folder", [
            "Click 'Browse' next to Data File (.xlsx) and pick your Excel survey file.",
            "Click 'Browse' next to Include.txt file and pick your include file (a text file listing the columns you want, one column name per line).",
            "Click 'Browse' next to Output folder and pick where you want the results saved."
        ]),
        ("(Optional) Enable advanced checks", [
            "All toggles are off by default. Enable only the checks you need:",
            "Automatically search the data file for open ends: Finds and adds open-ended/text columns to your output.",
            "Search file for specific words (words.txt): Flags and highlights cells containing words from a list you provide (one word per line in a .txt file). Click 'Browse' to select your word list.",
            "Check for duplicate postcode/YOB pairs: Flags and highlights rows where the combination of postcode and year of birth is duplicated.",
            "Enable length checks: Flags and highlights cells that are unusually long compared to others in the same column. You can set the multiplier (default is 10) to control what counts as 'unusually long' (e.g., 10× the median length for that column)."
        ]),
        ("Run the extraction", [
            "Click the Run button.",
            "The program will show you what it's doing (e.g., loading your file, running checks, creating output files).",
            "When finished, you'll see a message that your files are ready."
        ]),
        ("Check your results", [
            "Go to the output folder you chose. You'll find:",
            "An overall Excel file with your selected columns.",
            "Separate Excel files for each language in your data.",
            "A log file with details about what was done.",
            "(If you used open end detection) a new include file with open ends added.",
            "(If you enabled advanced checks) highlighted cells and a CHECKS column in the Excel output, showing which rows were flagged and why."
        ])
    ]
    for step, substeps in steps:
        doc.add_paragraph(step, style='List Number')
        for sub in substeps:
            doc.add_paragraph(sub, style='List Bullet 2')

    doc.add_heading('Tips', level=1)
    tips = [
        "If you see an error, check the error_log.txt file in your output folder for details.",
        "You can use the same include.txt file for different surveys—just update the column names as needed. Each column name should be on its own line (no semicolons).",
        "The program will never overwrite your original include.txt file."
    ]
    for tip in tips:
        doc.add_paragraph(tip, style='List Bullet')

    doc.add_heading('Need help?', level=1)
    doc.add_paragraph(
        "If you have questions or problems, contact your data team or open an issue on the GitHub page."
    )

    doc.add_paragraph("\nDataExtractQC makes survey data extraction fast and easy—no technical skills required!", style=None)

    doc.save(path)

if __name__ == "__main__":
    create_user_guide_docx("DataExtractQC_User_Guide.docx")
